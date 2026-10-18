#!/usr/bin/env python3
"""Build the downloadable Word templates offered by the template articles.

Writes blog/downloads/*.docx. Upload them to the site (for example
https://lynkk.ai/blog/downloads/<file>) and add the download link to the
matching article once the files are live.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "downloads"
GREY = RGBColor(0x59, 0x59, 0x59)


def new_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(72)
    return doc


def fields(doc, pairs):
    for label, value in pairs:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)
        p.paragraph_format.space_after = Pt(2)


def table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    doc.add_paragraph()


def footer(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Template by Lynkk (lynkk.ai). Lynkk can fill in notes, decisions and action items for you.")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY


def agenda(doc, title, meta, goal, rows, extra=None):
    doc.add_heading(title, level=1)
    fields(doc, meta)
    if goal:
        doc.add_heading("Goal", level=2)
        doc.add_paragraph(goal)
    table(doc, ["Time", "Topic", "Owner", "Decision needed"], rows)
    if extra:
        doc.add_heading(extra[0], level=2)
        doc.add_paragraph(extra[1])


def build_agendas():
    doc = new_doc()
    doc.add_heading("Meeting Agenda Templates", level=0)
    doc.add_paragraph("Replace the text in [brackets]. Delete the templates you do not need.")
    base = [("Meeting", "[ ]"), ("Date", "[ ]"), ("Time", "[ ]"),
            ("Location or link", "[ ]"), ("Facilitator", "[ ]"), ("Note-taker", "[ ]")]
    agenda(doc, "1. Simple meeting agenda", base, "[One sentence: what must be true when the meeting ends]",
           [["[10:00]", "Wins and updates", "All", "None"],
            ["[10:05]", "[Topic]", "[Name]", "[Decision]"],
            ["[10:15]", "[Topic]", "[Name]", "[Decision]"],
            ["[10:25]", "Recap and next steps", "[Name]", "Confirm actions"]],
           ("Pre-reading", "[Link]"))
    doc.add_page_break()
    agenda(doc, "2. Team meeting agenda", base[:4], None,
           [["5 min", "Check-in", "All", "None"],
            ["5 min", "Review last meeting's action items", "[Name]", "Close or reassign"],
            ["10 min", "Metrics and priorities", "[Name]", "None"],
            ["15 min", "Discussion: [topic]", "[Name]", "[Decision]"],
            ["5 min", "Action items and owners", "[Note-taker]", "Confirm owners and dates"]])
    doc.add_page_break()
    agenda(doc, "3. Client meeting agenda", [("Client", "[ ]")] + base[1:4],
           "[Agree on scope for phase 2]",
           [["5 min", "Introductions and goals", "Both teams", "None"],
            ["10 min", "Progress since last meeting", "[Name]", "None"],
            ["20 min", "Phase 2 scope and timeline", "[Name]", "Approve scope"],
            ["10 min", "Open questions from the client", "[Client]", "None"],
            ["5 min", "Next steps and owners", "[Name]", "Confirm actions"]])
    doc.add_page_break()
    agenda(doc, "4. One-on-one agenda", [("Manager", "[ ]"), ("Report", "[ ]"), ("Date", "[ ]")], None,
           [["5 min", "How are things going?", "Report", "None"],
            ["10 min", "Report's topics", "Report", "[ ]"],
            ["5 min", "Progress on goals", "Both", "None"],
            ["5 min", "Feedback both ways", "Both", "None"],
            ["5 min", "Action items", "Both", "Confirm actions"]])
    doc.add_page_break()
    agenda(doc, "5. Staff meeting agenda", base[:4], None,
           [["5 min", "Announcements", "[Manager]", "None"],
            ["15 min", "Department updates, 2 min each", "Department leads", "None"],
            ["15 min", "Discussion: [policy or process change]", "[Name]", "[Decision]"],
            ["5 min", "Questions from the team", "All", "None"],
            ["5 min", "Action items and owners", "[Note-taker]", "Confirm actions"]])
    doc.add_page_break()
    agenda(doc, "6. Weekly sales meeting agenda", base[:4], None,
           [["5 min", "Wins from last week", "Reps", "None"],
            ["10 min", "Pipeline and forecast vs target", "[Sales lead]", "None"],
            ["20 min", "Deal reviews: 2 to 3 deals that need help", "Deal owners", "Next step per deal"],
            ["5 min", "Enablement: one skill or message", "[Name]", "None"],
            ["5 min", "Actions and owners", "[Note-taker]", "Confirm actions"]])
    doc.add_page_break()
    agenda(doc, "7. All-hands agenda", base[:4], None,
           [["2 min", "Welcome and agenda", "[Host]", "None"],
            ["10 min", "Company update: goals and key metrics", "[CEO]", "None"],
            ["15 min", "Team spotlights, 2 teams", "[Leads]", "None"],
            ["15 min", "Topic deep dive: [ ]", "[Name]", "None"],
            ["15 min", "Q&A, including pre-submitted questions", "[Host]", "None"],
            ["3 min", "Close and next steps", "[Host]", "None"]])
    doc.add_page_break()
    doc.add_heading("8. Committee or formal meeting agenda", level=1)
    fields(doc, [("Committee", "[ ]"), ("Date", "[ ]"), ("Time", "[ ]"), ("Location", "[ ]"), ("Chair", "[ ]")])
    for item in ["1. Call to order", "2. Attendance and quorum", "3. Approval of the agenda",
                 "4. Approval of minutes of [date]", "5. Reports: chair, treasurer, committees",
                 "6. Old business: [item]", "7. New business: [item] (motion expected: [ ])",
                 "8. Next meeting date", "9. Adjournment"]:
        doc.add_paragraph(item)
    footer(doc)
    doc.save(OUT / "meeting-agenda-templates.docx")


def build_minutes():
    doc = new_doc()
    doc.add_heading("Minutes of Meeting", level=0)
    fields(doc, [("Organisation", "[Company / Department]"), ("Meeting", "[ ]"), ("MoM No.", "[ ]"),
                 ("Date", "[ ]"), ("Time", "[ ] to [ ]"), ("Venue", "[Room / video link]"),
                 ("Chairperson", "[ ]"), ("Minutes recorded by", "[ ]")])
    doc.add_heading("Attendance", level=2)
    fields(doc, [("Present", "[Name, designation]"), ("Absent", "[ ]"), ("Leave of absence", "[ ]")])
    doc.add_heading("Agenda", level=2)
    for item in ["1. Confirmation of previous minutes", "2. [Topic]", "3. [Topic]", "4. Any other business"]:
        doc.add_paragraph(item)
    doc.add_heading("Discussion and decisions", level=2)
    table(doc, ["#", "Agenda item", "Discussion (2 to 4 lines)", "Decision"],
          [["1", "Previous minutes", "Minutes of [date] were read.", "Approved"],
           ["2", "[Topic]", "[ ]", "[ ]"], ["3", "[Topic]", "[ ]", "[ ]"]])
    doc.add_heading("Action items", level=2)
    table(doc, ["#", "Action", "Owner", "Deadline"],
          [["1", "[ ]", "[ ]", "[ ]"], ["2", "[ ]", "[ ]", "[ ]"], ["3", "[ ]", "[ ]", "[ ]"]])
    doc.add_heading("Next meeting", level=2)
    fields(doc, [("Date", "[ ]"), ("Time", "[ ]"), ("Venue", "[ ]")])
    doc.add_paragraph()
    fields(doc, [("Prepared by", "[ ]"), ("Approved by", "[Chairperson]"), ("Date", "[ ]")])
    footer(doc)
    doc.save(OUT / "minutes-of-meeting-format.docx")


def build_notes():
    doc = new_doc()
    doc.add_heading("Meeting Notes", level=0)
    fields(doc, [("Meeting", "[ ]"), ("Date", "[ ]"), ("Time", "[ ]"), ("Location or link", "[ ]"),
                 ("Facilitator", "[ ]"), ("Note-taker", "[ ]"), ("Attendees", "[ ]"), ("Absent", "[ ]")])
    doc.add_heading("Decisions", level=2)
    for _ in range(3):
        doc.add_paragraph("[Decision]", style="List Bullet")
    doc.add_heading("Action items", level=2)
    table(doc, ["Task", "Owner", "Due date"], [["[ ]", "[ ]", "[ ]"] for _ in range(4)])
    doc.add_heading("Key points by topic", level=2)
    for i in range(1, 4):
        doc.add_paragraph(f"{i}. [Topic]: [1 to 3 lines]")
    doc.add_heading("Open questions", level=2)
    doc.add_paragraph("[Question] | [Who will find out]", style="List Bullet")
    doc.add_heading("Next meeting", level=2)
    doc.add_paragraph("[Date and time]")
    footer(doc)
    doc.save(OUT / "meeting-notes-template.docx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    build_agendas()
    build_minutes()
    build_notes()
    for f in sorted(OUT.glob("*.docx")):
        print(f"wrote downloads/{f.name} ({f.stat().st_size // 1024} KB)")
